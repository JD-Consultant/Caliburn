"""把 `ExportDocument` 渲染成唯一官方公版 XLSX。"""

from __future__ import annotations

import io
from math import ceil

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet.worksheet import Worksheet

from .export import ExportDocument, ExportOpksEntry, ExportTaskEntry


XLSX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
SHEET_NAME = "職能基準表"
TABLE_HEADER_ROW = 12
ISSUED_BY_ICAP = "（iCAP 計畫執行單位提供）"
TABLE_HEADERS = (
    "主要職責",
    "工作任務",
    "工作產出",
    "行為指標",
    "職能級別",
    "職能內涵（K=knowledge知識）",
    "職能內涵（S=skills技能）",
)
ATTITUDE_HEADER = "職能內涵（A=attitude態度）"
NOTES_HEADER = "說明與補充事項"

_THIN = Side(style="thin", color="000000")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_BOLD = Font(bold=True)
_WRAP_TOP = Alignment(vertical="top", wrap_text=True)
_WRAP_CENTER = Alignment(vertical="top", horizontal="center", wrap_text=True)


def _put_text(
    sheet: Worksheet,
    row: int,
    column: int,
    value: str | None,
    *,
    bold: bool = False,
    centered: bool = False,
) -> Cell:
    """以明確的 string cell 寫入所有員工／模型／模板文字。"""

    cell = sheet.cell(row=row, column=column)
    if value is not None:
        cell.value = str(value)
        cell.data_type = "s"
    cell.alignment = _WRAP_CENTER if centered else _WRAP_TOP
    if bold:
        cell.font = _BOLD
    return cell


def _put_number(
    sheet: Worksheet,
    row: int,
    column: int,
    value: int | None,
) -> Cell:
    cell = sheet.cell(row=row, column=column)
    if value is not None:
        cell.value = value
    cell.alignment = _WRAP_CENTER
    return cell


def _merge(sheet: Worksheet, row: int, first: int, last: int) -> None:
    if last > first:
        sheet.merge_cells(
            start_row=row,
            start_column=first,
            end_row=row,
            end_column=last,
        )


def _stack(entries: tuple[ExportOpksEntry, ...]) -> str | None:
    values = tuple(
        f"{entry.position_code}{entry.text}"
        if entry.position_code is not None
        else entry.text
        for entry in entries
    )
    return "\n".join(values) or None


def _write_header(sheet: Worksheet, document: ExportDocument) -> None:
    header = document.header
    _put_text(
        sheet,
        1,
        1,
        f"{header.competency_name or document.title}職能基準",
        bold=True,
    )
    _merge(sheet, 1, 1, 7)

    _put_text(sheet, 3, 1, "職能基準代碼", bold=True)
    _put_text(sheet, 3, 2, ISSUED_BY_ICAP)
    _merge(sheet, 3, 2, 7)

    _put_text(sheet, 4, 1, "職能基準名稱\n（擇一填寫）", bold=True)
    _put_text(sheet, 4, 2, "職類", bold=True)
    _merge(sheet, 4, 3, 7)
    _put_text(sheet, 5, 2, "職業", bold=True)
    _put_text(sheet, 5, 3, header.competency_name)
    _merge(sheet, 5, 3, 7)
    sheet.merge_cells(start_row=4, start_column=1, end_row=5, end_column=1)

    _put_text(sheet, 6, 1, "所屬類別", bold=True)
    categories = (
        ("職類別", header.occupation_category_name, "職類別代碼", ISSUED_BY_ICAP),
        ("職業別", header.occupation_name, "職業別代碼", header.occupation_code),
        ("行業別", header.industry_name, "行業別代碼", header.industry_code),
    )
    for row, (label, name, code_label, code) in enumerate(categories, start=6):
        _put_text(sheet, row, 2, label, bold=True)
        _put_text(sheet, row, 3, name)
        _put_text(sheet, row, 5, code_label, bold=True)
        _put_text(sheet, row, 6, code)
        _merge(sheet, row, 3, 4)
        _merge(sheet, row, 6, 7)
    sheet.merge_cells(start_row=6, start_column=1, end_row=8, end_column=1)

    _put_text(sheet, 9, 1, "工作描述", bold=True)
    _put_text(sheet, 9, 2, header.work_description)
    _merge(sheet, 9, 2, 7)
    _put_text(sheet, 10, 1, "基準級別", bold=True)
    _put_number(sheet, 10, 2, header.competency_level)
    _merge(sheet, 10, 2, 7)


def _write_task_row(sheet: Worksheet, row: int, entry: ExportTaskEntry) -> None:
    task_text = (
        f"{entry.position_code}{entry.statement}"
        if entry.position_code is not None
        else entry.statement
    )
    _put_text(sheet, row, 2, task_text)
    _put_text(sheet, row, 3, _stack(entry.outputs))
    _put_text(sheet, row, 4, _stack(entry.indicators))
    _put_number(sheet, row, 5, entry.competency_level)
    _put_text(sheet, row, 6, _stack(entry.knowledge))
    _put_text(sheet, row, 7, _stack(entry.skills))


def _write_table(sheet: Worksheet, document: ExportDocument) -> int:
    for column, title in enumerate(TABLE_HEADERS, start=1):
        _put_text(sheet, TABLE_HEADER_ROW, column, title, bold=True, centered=True)

    row = TABLE_HEADER_ROW + 1
    for section in document.duties:
        if not section.tasks:
            _put_text(sheet, row, 1, f"{section.position_code}{section.statement}")
            row += 1
            continue
        first = row
        _put_text(sheet, row, 1, f"{section.position_code}{section.statement}")
        for entry in section.tasks:
            _write_task_row(sheet, row, entry)
            row += 1
        if row - first > 1:
            sheet.merge_cells(
                start_row=first,
                start_column=1,
                end_row=row - 1,
                end_column=1,
            )

    # 未分組 Task 仍是官方 Task 列，Duty 欄刻意保持空白。
    for entry in document.unassigned_tasks:
        _write_task_row(sheet, row, entry)
        row += 1
    return row


def _write_block(sheet: Worksheet, row: int, title: str, body: str | None) -> int:
    _put_text(sheet, row, 1, title, bold=True)
    _merge(sheet, row, 1, 7)
    _put_text(sheet, row + 1, 1, body)
    _merge(sheet, row + 1, 1, 7)
    return row + 3


def _estimated_height(sheet: Worksheet, row: int) -> float:
    longest = max(
        (
            len(str(cell.value).split("\n", 1)[0])
            for cell in sheet[row]
            if cell.value is not None
        ),
        default=1,
    )
    return max(24.0, min(180.0, 18.0 * ceil(longest / 24)))


def _apply_layout(sheet: Worksheet, last_row: int) -> None:
    widths = {"A": 22, "B": 24, "C": 42, "D": 42, "E": 8, "F": 30, "G": 30}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for row in range(1, last_row + 1):
        sheet.row_dimensions[row].height = _estimated_height(sheet, row)
        for column in range(1, 8):
            cell = sheet.cell(row=row, column=column)
            cell.border = _BORDER
            if cell.alignment == Alignment():
                cell.alignment = _WRAP_TOP

    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = sheet.ORIENTATION_LANDSCAPE
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = f"{TABLE_HEADER_ROW}:{TABLE_HEADER_ROW}"
    sheet.print_area = f"A1:G{last_row}"


def render_xlsx(document: ExportDocument) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = SHEET_NAME
    _write_header(sheet, document)
    row = _write_table(sheet, document)
    row = _write_block(sheet, row, ATTITUDE_HEADER, _stack(document.attitudes))
    row = _write_block(sheet, row, NOTES_HEADER, document.header.notes)
    _apply_layout(sheet, row - 1)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


__all__ = ["XLSX_MEDIA_TYPE", "SHEET_NAME", "render_xlsx"]
