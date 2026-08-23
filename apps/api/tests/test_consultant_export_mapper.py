from io import BytesIO
from uuid import uuid4

from openpyxl import load_workbook

from app.adapters.xlsx import render_xlsx
from app.consultant.state import (
    ApprovedDuty,
    ApprovedJobDocument,
    ApprovedOpksItem,
    ApprovedOpksKind,
    ApprovedTask,
)
from app.export import assemble_approved_export_document


def test_new_approved_document_projects_to_the_retained_xlsx_shape() -> None:
    document_id = uuid4()
    source_id = uuid4()
    duty_id = uuid4()
    assigned_id = uuid4()
    second_assigned_id = uuid4()
    orphan_id = uuid4()
    indicator_id = uuid4()
    document = ApprovedJobDocument(
        document_id=document_id,
        job_title="採購專員",
        occupation_category_name="企業支援",
        occupation_name="採購人員",
        occupation_code="EMPLOYEE-CLASSIFICATION",
        industry_name="製造業",
        industry_code="EMPLOYEE-INDUSTRY",
        competency_level=3,
        duties=(
            ApprovedDuty(duty_id=duty_id, statement="採購作業", display_order=0),
        ),
        tasks=(
            ApprovedTask(
                task_id=assigned_id,
                duty_id=duty_id,
                statement="整理採購需求",
                action="整理",
                object="採購需求",
                display_order=0,
                competency_level=2,
            ),
            ApprovedTask(
                task_id=second_assigned_id,
                duty_id=duty_id,
                statement="覆核採購需求",
                action="覆核",
                object="採購需求",
                display_order=1,
                competency_level=3,
            ),
            ApprovedTask(
                task_id=orphan_id,
                duty_id=None,
                statement="追蹤緊急缺料",
                action="追蹤",
                object="緊急缺料",
                display_order=2,
            ),
        ),
        opks=(
            ApprovedOpksItem(
                item_id=uuid4(),
                kind=ApprovedOpksKind.OUTPUT,
                text="採購需求清單",
                display_order=0,
                task_ids=(assigned_id,),
                evidence_source_ids=(source_id,),
            ),
            ApprovedOpksItem(
                item_id=indicator_id,
                kind=ApprovedOpksKind.PERFORMANCE_INDICATOR,
                text="需求內容完整",
                display_order=0,
                task_ids=(assigned_id,),
                evidence_source_ids=(source_id,),
            ),
            ApprovedOpksItem(
                item_id=uuid4(),
                kind=ApprovedOpksKind.KNOWLEDGE,
                text="採購流程知識",
                display_order=0,
                task_ids=(second_assigned_id,),
                indicator_ids=(indicator_id,),
                evidence_source_ids=(source_id,),
            ),
            ApprovedOpksItem(
                item_id=uuid4(),
                kind=ApprovedOpksKind.SKILL,
                text="需求彙整技能",
                display_order=0,
                task_ids=(assigned_id,),
                evidence_source_ids=(source_id,),
            ),
            ApprovedOpksItem(
                item_id=uuid4(),
                kind=ApprovedOpksKind.ATTITUDE,
                text="謹慎",
                display_order=0,
                evidence_source_ids=(source_id,),
            ),
        ),
    )

    exported = assemble_approved_export_document(document, title="採購職務")

    first_task, second_task = exported.duties[0].tasks
    assert exported.header.competency_name == "採購專員"
    assert exported.header.occupation_category_name == "企業支援"
    assert exported.header.occupation_name == "採購人員"
    assert exported.header.industry_name == "製造業"
    assert exported.header.occupation_code is None
    assert exported.header.industry_code is None
    assert first_task.position_code == "T1.1"
    assert first_task.competency_level == 2
    assert first_task.outputs[0].position_code == "O1.1.1"
    assert first_task.indicators[0].position_code == "P1.1.1"
    assert second_task.outputs == ()
    assert second_task.indicators == ()
    assert first_task.knowledge[0].position_code == "K01"
    assert second_task.knowledge[0].position_code == "K01"
    assert first_task.skills[0].position_code == "S01"
    assert second_task.skills == ()
    assert exported.unassigned_tasks[0].task_id == str(orphan_id)
    assert exported.unassigned_tasks[0].position_code is None
    assert exported.attitudes[0].position_code == "A01"

    workbook = load_workbook(BytesIO(render_xlsx(exported)), data_only=False)
    sheet = workbook.active
    assert sheet is not None
    assert sheet["B3"].value is None
    assert sheet["F6"].value is None
    assert sheet["F7"].value is None
    assert sheet["F8"].value is None
    assert sheet["B10"].value == 3
    assert any(cell.value == "A01謹慎" for row in sheet.iter_rows() for cell in row)
