"""Real PostgreSQL + HTTP vertical for the local job-analysis workspace."""

from __future__ import annotations

from io import BytesIO

import httpx
import pytest
import pytest_asyncio
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.adapters.job_analysis_postgres import SqlAlchemyJobAnalysisUnitOfWork
from app.adapters.job_analysis_postgres.models import JobAnalysisJournalRow
from app.api.job_analysis_deps import (
    get_job_analysis_adapter,
    get_job_analysis_uow_factory,
)
from app.job_analysis.application import load_document
from app.core.domain import TaskFields
from app.job_analysis.llm import (
    IdentityRelation,
    SignalDisposition,
    TaskAnalysisWire,
    WireAnchor,
    WireNextQuestion,
    WireSignal,
    WireTaskChange,
    WireTaskFields,
)
from app.job_analysis.providers import (
    OpenRouterAdapter,
    OpenRouterConfig,
    TransportResponse,
)
from app.main import app


pytestmark = pytest.mark.asyncio


class _ConsultantTransport:
    def __init__(self) -> None:
        self.calls = 0

    async def __call__(self, *, url, headers, body, timeout):
        self.calls += 1
        result = TaskAnalysisWire(
            work_signals=(
                WireSignal(
                    anchors=(
                        WireAnchor(
                            turn_ordinal=2,
                            quote="我每週彙整營運週報",
                        ),
                    ),
                    relation=IdentityRelation.NO_MATCH,
                    disposition=SignalDisposition.TASK_CHANGE,
                    change=WireTaskChange.ADD,
                    task=WireTaskFields(
                        statement="每週彙整營運週報",
                        action="彙整",
                        object="營運週報",
                        purpose_result="讓主管掌握營運狀況",
                    ),
                ),
            ),
            next_question=WireNextQuestion(
                text="這份週報通常提供給誰？",
            ),
        )
        return TransportResponse(
            status_code=200,
            body={
                "model": "anthropic/claude-opus-5",
                "choices": [
                    {
                        "message": {
                            "content": result.model_dump_json(),
                            "refusal": None,
                        },
                        "finish_reason": "stop",
                    }
                ],
            },
        )


@pytest_asyncio.fixture
async def postgres_api_client(postgres_session_factory):
    factory = lambda: SqlAlchemyJobAnalysisUnitOfWork(postgres_session_factory)
    provider = _ConsultantTransport()
    adapter = OpenRouterAdapter(
        config=OpenRouterConfig(
            model="anthropic/claude-opus-5",
            provider_order=("anthropic",),
            max_output_tokens=4096,
            timeout_seconds=90,
        ),
        api_key="sk-test",
        transport=provider,
    )
    app.dependency_overrides[get_job_analysis_uow_factory] = lambda: factory
    app.dependency_overrides[get_job_analysis_adapter] = lambda: adapter
    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(
        transport=transport,
        base_url="http://test",
    ) as client:
        yield client, factory, provider
    app.dependency_overrides.clear()


def task_payload(
    statement: str,
    *,
    purpose_result: str | None = None,
    frequency_text: str | None = None,
    duty_id: str | None = None,
    competency_level: int | None = None,
):
    return {
        "statement": statement,
        "purpose_result": purpose_result,
        "context": None,
        "frequency_text": frequency_text,
        "responsibility_role": None,
        "enablers": [],
        "duty_id": duty_id,
        "competency_level": competency_level,
    }


async def test_current_jd_duty_routes_author_and_unassign_tasks_on_delete(
    postgres_api_client,
    cleanup_job_analysis_rows,
):
    client, _, _ = postgres_api_client
    document_id = cleanup_job_analysis_rows
    root = f"/api/v1/job-analysis/documents/{document_id}"

    assert (await client.put(root, json={"title": "門市營運專員"})).status_code == 201

    duty = await client.post(
        f"{root}/duties",
        headers={"Idempotency-Key": "duty-add-1"},
        json={"statement": "門市營運"},
    )
    assert duty.status_code == 201
    assert duty.json()["duty_id"] == "duty-add-1-d0"

    task = await client.post(
        f"{root}/tasks",
        headers={"Idempotency-Key": "duty-task-1"},
        json=task_payload(
            "盤點門市耗材",
            duty_id=duty.json()["duty_id"],
            competency_level=4,
        ),
    )
    assert task.status_code == 201
    task_id = task.json()["task_id"]
    assert task.json()["duty_id"] == duty.json()["duty_id"]
    assert task.json()["competency_level"] == 4

    renamed = await client.put(
        f"{root}/duties/{duty.json()['duty_id']}",
        headers={"Idempotency-Key": "duty-edit-1"},
        json={"statement": "門市日常營運"},
    )
    assert renamed.status_code == 200
    assert renamed.json()["display_order"] == 0

    reordered = await client.put(
        f"{root}/duty-order",
        headers={"Idempotency-Key": "duty-order-1"},
        json={"ordered_duty_ids": [duty.json()["duty_id"]]},
    )
    assert reordered.status_code == 200

    deleted = await client.delete(
        f"{root}/duties/{duty.json()['duty_id']}",
        headers={"Idempotency-Key": "duty-delete-1"},
    )
    reloaded = await client.get(root)

    assert deleted.status_code == 204
    assert reloaded.status_code == 200
    assert reloaded.json()["tasks"] == [
        {
            **task.json(),
            "duty_id": None,
        }
    ]
    assert reloaded.json()["tasks"][0]["task_id"] == task_id


async def test_current_jd_duty_routes_reject_missing_duty_and_invalid_order(
    postgres_api_client,
    cleanup_job_analysis_rows,
):
    client, _, _ = postgres_api_client
    document_id = cleanup_job_analysis_rows
    root = f"/api/v1/job-analysis/documents/{document_id}"

    assert (await client.put(root, json={"title": "門市營運專員"})).status_code == 201
    missing = await client.put(
        f"{root}/duties/missing-duty",
        headers={"Idempotency-Key": "duty-missing-1"},
        json={"statement": "不存在"},
    )
    invalid_order = await client.put(
        f"{root}/duty-order",
        headers={"Idempotency-Key": "duty-order-invalid-1"},
        json={"ordered_duty_ids": ["missing-duty"]},
    )

    assert missing.status_code == 404
    assert missing.json()["type"].endswith("/duty-not-found")
    assert invalid_order.status_code == 422
    assert invalid_order.json()["type"].endswith("/invalid-duty-order")


async def test_local_web_task_editing_survives_reload_and_rename_is_metadata_only(
    postgres_api_client,
    postgres_session_factory,
    cleanup_job_analysis_rows,
):
    client, factory, _ = postgres_api_client
    document_id = cleanup_job_analysis_rows
    root = f"/api/v1/job-analysis/documents/{document_id}"

    assert (await client.put(root, json={"title": "門市營運專員"})).status_code == 201
    first = await client.post(
        f"{root}/tasks",
        headers={"Idempotency-Key": "vertical-add-1"},
        json=task_payload("盤點門市耗材"),
    )
    assert first.status_code == 201
    first_id = first.json()["task_id"]
    assert first.json()["purpose_result"] is None

    edited = await client.put(
        f"{root}/tasks/{first_id}",
        headers={"Idempotency-Key": "vertical-edit-1"},
        json=task_payload(
            "盤點門市耗材",
            purpose_result="確保安全庫存充足",
            frequency_text="每月一次",
        ),
    )
    assert edited.status_code == 200

    second = await client.post(
        f"{root}/tasks",
        headers={"Idempotency-Key": "vertical-add-2"},
        json=task_payload("每週彙整營運週報"),
    )
    second_id = second.json()["task_id"]
    reordered = await client.put(
        f"{root}/task-order",
        headers={"Idempotency-Key": "vertical-order-1"},
        json={"ordered_task_ids": [second_id, first_id]},
    )
    assert [item["task_id"] for item in reordered.json()] == [second_id, first_id]

    before_rename = await load_document(factory, document_id)
    assert before_rename is not None
    async with postgres_session_factory() as session:
        journal_before = await session.scalar(
            select(func.count())
            .select_from(JobAnalysisJournalRow)
            .where(JobAnalysisJournalRow.document_id == document_id)
        )

    renamed = await client.put(root, json={"title": "資深門市營運專員"})
    reloaded = await client.get(root)
    after_rename = await load_document(factory, document_id)
    async with postgres_session_factory() as session:
        journal_after = await session.scalar(
            select(func.count())
            .select_from(JobAnalysisJournalRow)
            .where(JobAnalysisJournalRow.document_id == document_id)
        )

    assert renamed.status_code == 200
    assert reloaded.status_code == 200
    assert reloaded.json()["title"] == "資深門市營運專員"
    assert [item["task_id"] for item in reloaded.json()["tasks"]] == [
        second_id,
        first_id,
    ]
    assert reloaded.json()["tasks"][1]["purpose_result"] == "確保安全庫存充足"
    assert reloaded.json()["tasks"][1]["frequency_text"] == "每月一次"
    assert after_rename is not None
    assert after_rename.document.authority_generation == (
        before_rename.document.authority_generation
    )
    assert after_rename.state == before_rename.state
    assert journal_after == journal_before


async def test_put_jd_header_trims_optional_text_and_returns_server_readiness(
    postgres_api_client,
    cleanup_job_analysis_rows,
):
    client, factory, _ = postgres_api_client
    document_id = cleanup_job_analysis_rows
    root = f"/api/v1/job-analysis/documents/{document_id}"
    assert (await client.put(root, json={"title": "門市營運專員"})).status_code == 201

    put = await client.put(
        f"{root}/jd-header",
        headers={"Idempotency-Key": "header-1"},
        json={
            "competency_name": " 門市營運專員 ",
            "notes": "   ",
            "competency_level": 4,
        },
    )
    reloaded = await client.get(root)
    replay = await client.put(
        f"{root}/jd-header",
        headers={"Idempotency-Key": "header-1"},
        json={
            "competency_name": " 門市營運專員 ",
            "notes": "   ",
            "competency_level": 4,
        },
    )
    persisted = await load_document(factory, document_id)

    assert put.status_code == 200
    assert put.json()["competency_name"] == "門市營運專員"
    assert put.json()["notes"] is None
    assert put.json()["occupation_name"] is None
    assert reloaded.status_code == 200
    assert reloaded.json()["jd_header"] == put.json()
    assert reloaded.json()["readiness"]["issues"] == [
        {"code": "work_description_missing"}
    ]
    assert replay.status_code == 200
    assert replay.json() == put.json()
    assert persisted is not None
    assert persisted.document.authority_generation == 1


async def test_export_returns_one_sheet_xlsx_without_mutating_authority_or_journal(
    postgres_api_client,
    postgres_session_factory,
    cleanup_job_analysis_rows,
):
    client, factory, _ = postgres_api_client
    document_id = cleanup_job_analysis_rows
    root = f"/api/v1/job-analysis/documents/{document_id}"
    assert (await client.put(root, json={"title": "門市營運專員"})).status_code == 201

    before = await load_document(factory, document_id)
    assert before is not None
    async with postgres_session_factory() as session:
        journal_before = await session.scalar(
            select(func.count())
            .select_from(JobAnalysisJournalRow)
            .where(JobAnalysisJournalRow.document_id == document_id)
        )

    exported = await client.get(f"{root}/export")
    after = await load_document(factory, document_id)
    async with postgres_session_factory() as session:
        journal_after = await session.scalar(
            select(func.count())
            .select_from(JobAnalysisJournalRow)
            .where(JobAnalysisJournalRow.document_id == document_id)
        )

    assert exported.status_code == 200
    assert exported.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "filename*=UTF-8''" in exported.headers["content-disposition"]
    assert "%E9%96%80%E5%B8%82%E7%87%9F%E9%81%8B%E5%B0%88%E5%93%A1.xlsx" in (
        exported.headers["content-disposition"]
    )
    workbook = load_workbook(BytesIO(exported.content), read_only=True)
    assert workbook.sheetnames == ["職能基準表"]
    workbook.close()
    assert after is not None
    assert after.document.authority_generation == before.document.authority_generation
    assert after.state == before.state
    assert journal_after == journal_before


async def test_export_returns_document_not_found_problem(
    postgres_api_client,
):
    client, _, _ = postgres_api_client

    response = await client.get(
        "/api/v1/job-analysis/documents/00000000-0000-0000-0000-000000000099/export"
    )

    assert response.status_code == 404
    assert response.headers["content-type"].startswith("application/problem+json")
    assert response.json()["type"].endswith("/document-not-found")


async def test_put_jd_header_is_a_full_replacement_that_nulls_omitted_fields(
    postgres_api_client,
    cleanup_job_analysis_rows,
):
    client, factory, _ = postgres_api_client
    document_id = cleanup_job_analysis_rows
    root = f"/api/v1/job-analysis/documents/{document_id}"
    assert (await client.put(root, json={"title": "門市營運專員"})).status_code == 201

    populated = await client.put(
        f"{root}/jd-header",
        headers={"Idempotency-Key": "header-populate"},
        json={
            "competency_name": "門市營運專員",
            "occupation_name": "零售服務人員",
            "work_description": "負責門市日常營運。",
            "competency_level": 3,
            "notes": "初版備註",
        },
    )
    before = await load_document(factory, document_id)

    replaced = await client.put(
        f"{root}/jd-header",
        headers={"Idempotency-Key": "header-replace"},
        json={
            "competency_name": "資深門市營運專員",
            "work_description": "負責門市日常營運與人員帶領。",
            "competency_level": 4,
        },
    )
    reloaded = await client.get(root)
    after = await load_document(factory, document_id)

    assert populated.status_code == 200
    assert before is not None
    assert replaced.status_code == 200
    assert replaced.json()["occupation_name"] is None
    assert replaced.json()["notes"] is None
    assert reloaded.status_code == 200
    assert reloaded.json()["jd_header"] == replaced.json()
    assert reloaded.json()["jd_header"]["occupation_name"] is None
    assert after is not None
    assert after.state.jd_header.occupation_name is None
    assert after.document.authority_generation == (
        before.document.authority_generation + 1
    )


async def test_put_jd_header_rejects_an_idempotency_collision_as_a_problem(
    postgres_api_client,
    cleanup_job_analysis_rows,
):
    client, _, _ = postgres_api_client
    document_id = cleanup_job_analysis_rows
    root = f"/api/v1/job-analysis/documents/{document_id}"
    assert (await client.put(root, json={"title": "門市營運專員"})).status_code == 201

    first = await client.put(
        f"{root}/jd-header",
        headers={"Idempotency-Key": "header-collision"},
        json={"competency_name": "門市營運專員"},
    )
    collision = await client.put(
        f"{root}/jd-header",
        headers={"Idempotency-Key": "header-collision"},
        json={"competency_name": "資深門市營運專員"},
    )

    assert first.status_code == 200
    assert collision.status_code == 409
    assert collision.headers["content-type"] == "application/problem+json"
    assert collision.json()["type"] == (
        "https://caliburn.dev/problems/job-analysis/idempotency-conflict"
    )


async def test_put_jd_header_reports_no_op_and_invalid_bodies_as_typed_422(
    postgres_api_client,
    cleanup_job_analysis_rows,
):
    client, factory, _ = postgres_api_client
    document_id = cleanup_job_analysis_rows
    root = f"/api/v1/job-analysis/documents/{document_id}"
    assert (await client.put(root, json={"title": "門市營運專員"})).status_code == 201

    no_op = await client.put(
        f"{root}/jd-header",
        headers={"Idempotency-Key": "header-no-op"},
        json={},
    )
    invalid_level = await client.put(
        f"{root}/jd-header",
        headers={"Idempotency-Key": "header-invalid-level"},
        json={"competency_level": 7},
    )
    invalid_body = await client.put(
        f"{root}/jd-header",
        headers={"Idempotency-Key": "header-invalid-body"},
        json={"unknown_header_field": "nope"},
    )
    missing_key = await client.put(
        f"{root}/jd-header",
        json={"competency_name": "門市營運專員"},
    )
    persisted = await load_document(factory, document_id)

    for response in (no_op, invalid_level, invalid_body, missing_key):
        assert response.status_code == 422
        assert response.headers["content-type"] == "application/problem+json"
        assert response.json()["type"] == (
            "https://caliburn.dev/problems/job-analysis/invalid-request"
        )
    assert persisted is not None
    assert persisted.document.authority_generation == 0


@pytest.mark.parametrize("invalid_level", ["4", True])
async def test_put_jd_header_rejects_non_integer_levels_as_typed_422(
    postgres_api_client,
    cleanup_job_analysis_rows,
    invalid_level,
):
    client, factory, _ = postgres_api_client
    document_id = cleanup_job_analysis_rows
    root = f"/api/v1/job-analysis/documents/{document_id}"
    assert (await client.put(root, json={"title": "門市營運專員"})).status_code == 201

    response = await client.put(
        f"{root}/jd-header",
        headers={"Idempotency-Key": "header-invalid-type"},
        json={"competency_level": invalid_level},
    )
    persisted = await load_document(factory, document_id)

    assert response.status_code == 422
    assert response.headers["content-type"] == "application/problem+json"
    assert response.json()["type"] == (
        "https://caliburn.dev/problems/job-analysis/invalid-request"
    )
    assert persisted is not None
    assert persisted.document.authority_generation == 0


async def test_consultant_turn_proposal_decision_and_reload_use_real_postgres(
    postgres_api_client,
    postgres_session_factory,
    cleanup_job_analysis_rows,
):
    client, factory, provider = postgres_api_client
    document_id = cleanup_job_analysis_rows
    root = f"/api/v1/job-analysis/documents/{document_id}"
    await client.put(root, json={"title": "門市營運專員"})

    first = await client.post(
        f"{root}/turns",
        headers={"Idempotency-Key": "pg-turn-1"},
        json={"text": "我每週彙整營運週報"},
    )
    replay = await client.post(
        f"{root}/turns",
        headers={"Idempotency-Key": "pg-turn-1"},
        json={"text": "我每週彙整營運週報"},
    )
    proposal_id = first.json()["proposals"][0]["proposal_id"]
    accepted = await client.post(
        f"{root}/proposals/{proposal_id}/decisions",
        headers={"Idempotency-Key": "pg-decision-1"},
        json={"decision": "accepted"},
    )
    reloaded = await client.get(f"{root}/consultation")
    persisted = await load_document(factory, document_id)
    async with postgres_session_factory() as session:
        journal_count = await session.scalar(
            select(func.count())
            .select_from(JobAnalysisJournalRow)
            .where(JobAnalysisJournalRow.document_id == document_id)
        )

    assert first.status_code == 200
    assert replay.json() == first.json()
    assert provider.calls == 1
    assert accepted.status_code == 200
    assert accepted.json()["tasks"][0]["statement"] == "每週彙整營運週報"
    assert accepted.json()["active_question"]["text"] == "這份週報通常提供給誰？"
    assert reloaded.json() == accepted.json()
    assert persisted is not None
    assert len(persisted.conversation_turns) == 3
    assert journal_count == 3
