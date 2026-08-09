"""嚴格公版 XLSX 的版型、文字安全與列印設定。"""

from __future__ import annotations

import io

from openpyxl import load_workbook

from app.job_analysis.application import (
    ExportDocument,
    ExportDutySection,
    ExportOpksEntry,
    ExportTaskEntry,
)
from app.job_analysis.domain import JdHeader
from app.job_analysis.application.export_xlsx import render_xlsx


def _document(*, header: JdHeader | None = None) -> ExportDocument:
    task = ExportTaskEntry(
        task_id="t1",
        position_code="T1.1",
        statement="每週彙整營運週報",
        competency_level=4,
        outputs=(ExportOpksEntry(position_code="O1.1.1", text="營運週報"),),
        indicators=(ExportOpksEntry(position_code="P1.1.1", text="資料正確"),),
        knowledge=(ExportOpksEntry(position_code="K01", text="資料整理"),),
        skills=(ExportOpksEntry(position_code="S01", text="試算表操作"),),
    )
    return ExportDocument(
        title="門市營運專員",
        header=header or JdHeader(competency_name="門市營運專員"),
        duties=(
            ExportDutySection(
                position_code="T1",
                statement="維運門市營運系統",
                tasks=(task,),
            ),
        ),
        attitudes=(ExportOpksEntry(position_code="A01", text="細心"),),
    )


def _workbook(document: ExportDocument):
    return load_workbook(io.BytesIO(render_xlsx(document)), data_only=False)


def test_workbook_has_only_the_official_form_sheet():
    workbook = _workbook(_document())

    assert workbook.sheetnames == ["職能基準表"]


def test_formula_shaped_employee_text_stays_literal_text():
    workbook = _workbook(
        _document(header=JdHeader(work_description="=1+1"))
    )

    cell = next(
        cell
        for row in workbook.active.iter_rows()
        for cell in row
        if cell.value == "=1+1"
    )
    assert cell.value == "=1+1"
    assert cell.data_type == "s"


def test_public_form_has_official_seven_column_headers_and_icap_notes():
    sheet = _workbook(_document()).active
    values = [cell.value for cell in sheet[12]]
    all_text = "\n".join(
        str(cell.value)
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    )

    assert values == [
        "主要職責",
        "工作任務",
        "工作產出",
        "行為指標",
        "職能級別",
        "職能內涵（K=knowledge知識）",
        "職能內涵（S=skills技能）",
    ]
    assert "（iCAP 計畫執行單位提供）" in {
        cell.value for row in sheet.iter_rows() for cell in row
    }
    assert "iCAP 版型缺漏" not in all_text
    assert "不代表勞動部認證" not in all_text


def test_unassigned_task_stays_in_the_official_task_column_without_custom_label():
    document = _document().model_copy(
        update={
            "unassigned_tasks": (
                ExportTaskEntry(
                    task_id="t2",
                    position_code=None,
                    statement="盤點耗材",
                    competency_level=None,
                ),
            )
        }
    )
    sheet = _workbook(document).active
    all_values = [cell.value for row in sheet.iter_rows() for cell in row]

    assert "盤點耗材" in all_values
    assert "（尚未歸入主要職責）" not in all_values
    assert "（尚未連結工作任務）" not in all_values
    task_row = next(row for row in sheet.iter_rows() if any(cell.value == "盤點耗材" for cell in row))
    assert task_row[0].value is None


def test_long_text_is_wrapped_and_empty_fields_remain_blank():
    long_text = "這是一段很長的中文工作描述，用來確認公版表格不會把員工文字截斷。" * 4
    document = _document(
        header=JdHeader(
            competency_name="門市營運專員",
            work_description=long_text,
            occupation_name=None,
            occupation_code=None,
        )
    )
    sheet = _workbook(document).active

    description_cell = next(
        cell
        for row in sheet.iter_rows()
        for cell in row
        if cell.value == long_text
    )
    assert description_cell.alignment.wrap_text is True
    occupation_code_label = next(
        cell for row in sheet.iter_rows() for cell in row if cell.value == "職業別代碼"
    )
    assert sheet.cell(occupation_code_label.row, occupation_code_label.column + 1).value is None


def test_shared_knowledge_and_skill_codes_are_written_as_same_codes():
    document = _document().model_copy(
        update={
            "duties": (
                ExportDutySection(
                    position_code="T1",
                    statement="維運門市營運系統",
                    tasks=(
                        _document().duties[0].tasks[0],
                        ExportTaskEntry(
                            task_id="t2",
                            position_code="T1.2",
                            statement="追蹤異常",
                            competency_level=3,
                            knowledge=(ExportOpksEntry(position_code="K01", text="資料整理"),),
                            skills=(ExportOpksEntry(position_code="S01", text="試算表操作"),),
                        ),
                    ),
                ),
            )
        }
    )
    sheet = _workbook(document).active
    stacked = [
        cell.value
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and "K01資料整理" in cell.value
    ]
    skills = [
        cell.value
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and "S01試算表操作" in cell.value
    ]

    assert len(stacked) == 2
    assert len(skills) == 2


def test_used_cells_have_borders_and_print_setup_is_single_page_wide():
    sheet = _workbook(_document()).active

    used_cells = [
        cell
        for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, min_col=1, max_col=7)
        for cell in row
        if cell.value is not None
    ]
    assert used_cells
    assert all(cell.border.left.style == "thin" for cell in used_cells)
    assert all(cell.border.right.style == "thin" for cell in used_cells)
    assert sheet.page_setup.orientation == "landscape"
    assert str(sheet.page_setup.paperSize) == sheet.PAPERSIZE_A4
    assert sheet.page_setup.fitToWidth == 1
    assert sheet.sheet_properties.pageSetUpPr.fitToPage is True
    assert sheet.print_title_rows == "$12:$12"
    assert sheet.print_area
