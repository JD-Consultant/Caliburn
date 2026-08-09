"""真 PostgreSQL 上驗證人工 authoring 到公版 XLSX 的完整垂直切片。"""

from __future__ import annotations

from io import BytesIO

import httpx
import pytest
import pytest_asyncio
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.adapters.job_analysis_postgres import SqlAlchemyJobAnalysisUnitOfWork
from app.adapters.job_analysis_postgres.models import JobAnalysisJournalRow
from app.api.job_analysis_deps import get_job_analysis_uow_factory
from app.job_analysis.application import load_document
from app.main import app


pytestmark = pytest.mark.asyncio


@pytest_asyncio.fixture
async def vertical_api_client(postgres_session_factory):
    factory = lambda: SqlAlchemyJobAnalysisUnitOfWork(postgres_session_factory)
    app.dependency_overrides[get_job_analysis_uow_factory] = lambda: factory
    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(
        transport=transport,
        base_url="http://test",
    ) as client:
        yield client, factory
    app.dependency_overrides.clear()


def task_payload(
    statement: str,
    *,
    duty_id: str | None = None,
    competency_level: int | None = None,
):
    return {
        "statement": statement,
        "purpose_result": None,
        "context": None,
        "frequency_text": None,
        "responsibility_role": None,
        "enablers": [],
        "duty_id": duty_id,
        "competency_level": competency_level,
    }


async def test_header_duty_task_opks_export_vertical(
    vertical_api_client,
    postgres_session_factory,
    cleanup_job_analysis_rows,
):
    client, factory = vertical_api_client
    document_id = cleanup_job_analysis_rows
    root = f"/api/v1/job-analysis/documents/{document_id}"

    assert (await client.put(root, json={"title": "門市營運專員"})).status_code == 201
    before = await load_document(factory, document_id)
    assert before is not None

    header = await client.put(
        f"{root}/jd-header",
        headers={"Idempotency-Key": "vertical-header"},
        json={
            "competency_name": "門市營運專員",
            "occupation_category_name": "零售服務",
            "occupation_name": "門市服務人員",
            "occupation_code": "=1+1",
            "industry_name": "零售業",
            "industry_code": "G47",
            "work_description": "負責門市營運與服務品質。",
            "competency_level": 4,
            "notes": "供內部職務說明使用。",
        },
    )
    assert header.status_code == 200

    first_duty = await client.post(
        f"{root}/duties",
        headers={"Idempotency-Key": "vertical-duty-first"},
        json={"statement": "門市日常營運"},
    )
    second_duty = await client.post(
        f"{root}/duties",
        headers={"Idempotency-Key": "vertical-duty-second"},
        json={"statement": "報表與盤點"},
    )
    assert first_duty.status_code == 201
    assert second_duty.status_code == 201
    first_duty_id = first_duty.json()["duty_id"]
    second_duty_id = second_duty.json()["duty_id"]

    first_task = await client.post(
        f"{root}/tasks",
        headers={"Idempotency-Key": "vertical-task-first"},
        json=task_payload(
            "盤點門市耗材",
            duty_id=first_duty_id,
            competency_level=4,
        ),
    )
    second_task = await client.post(
        f"{root}/tasks",
        headers={"Idempotency-Key": "vertical-task-second"},
        json=task_payload(
            "彙整營運週報",
            duty_id=second_duty_id,
        ),
    )
    assert first_task.status_code == 201
    assert second_task.status_code == 201
    first_task_id = first_task.json()["task_id"]
    second_task_id = second_task.json()["task_id"]

    reordered_duties = await client.put(
        f"{root}/duty-order",
        headers={"Idempotency-Key": "vertical-duty-order"},
        json={"ordered_duty_ids": [second_duty_id, first_duty_id]},
    )
    assert [duty["duty_id"] for duty in reordered_duties.json()] == [
        second_duty_id,
        first_duty_id,
    ]

    deleted_duty = await client.delete(
        f"{root}/duties/{second_duty_id}",
        headers={"Idempotency-Key": "vertical-duty-delete"},
    )
    assert deleted_duty.status_code == 204

    opks = []
    for key, kind, text, task_refs in (
        ("vertical-output-first", "output", "每日盤點", [first_task_id]),
        ("vertical-output-second", "output", "每週彙整", [first_task_id]),
        ("vertical-indicator", "indicator", "資料正確", [first_task_id]),
        ("vertical-knowledge", "knowledge", "=1+1", [first_task_id]),
        ("vertical-skill", "skill", "試算表操作", [first_task_id]),
        ("vertical-attitude", "attitude", "細心", []),
    ):
        response = await client.post(
            f"{root}/opks",
            headers={"Idempotency-Key": key},
            json={
                "entity_kind": kind,
                "text": text,
                "task_refs": task_refs,
                "indicator_refs": [],
            },
        )
        assert response.status_code == 201
        opks.append(response.json())

    output_ids = [
        item["entity_id"]
        for item in opks
        if item["entity_kind"] == "output"
    ]
    reordered_opks = await client.put(
        f"{root}/opks-order",
        headers={"Idempotency-Key": "vertical-opks-order"},
        json={
            "entity_kind": "output",
            "ordered_entity_ids": list(reversed(output_ids)),
        },
    )
    assert [item["entity_id"] for item in reordered_opks.json()] == list(
        reversed(output_ids)
    )

    reloaded = await client.get(root)
    assert reloaded.status_code == 200
    assert [task["task_id"] for task in reloaded.json()["tasks"]] == [
        first_task_id,
        second_task_id,
    ]
    assert reloaded.json()["tasks"][0]["competency_level"] == 4
    assert reloaded.json()["tasks"][1]["duty_id"] is None

    async with postgres_session_factory() as session:
        journal_rows = list(
            await session.scalars(
                select(JobAnalysisJournalRow)
                .where(JobAnalysisJournalRow.document_id == document_id)
                .order_by(JobAnalysisJournalRow.journal_sequence)
            )
        )
        journal_count = await session.scalar(
            select(func.count())
            .select_from(JobAnalysisJournalRow)
            .where(JobAnalysisJournalRow.document_id == document_id)
        )

    after = await load_document(factory, document_id)
    assert after is not None
    expected_entry_ids = {
        "consultant-opening",
        "vertical-header",
        "vertical-duty-first",
        "vertical-duty-second",
        "vertical-task-first",
        "vertical-task-second",
        "vertical-duty-order",
        "vertical-duty-delete",
        "vertical-output-first",
        "vertical-output-second",
        "vertical-indicator",
        "vertical-knowledge",
        "vertical-skill",
        "vertical-attitude",
        "vertical-opks-order",
    }
    assert journal_count == len(expected_entry_ids)
    assert {row.entry_id for row in journal_rows} == expected_entry_ids
    assert after.document.authority_generation == (
        before.document.authority_generation
        + len(expected_entry_ids - {"consultant-opening"})
    )

    exported = await client.get(f"{root}/export")
    assert exported.status_code == 200
    workbook = load_workbook(BytesIO(exported.content), data_only=False)
    assert workbook.sheetnames == ["職能基準表"]
    sheet = workbook.active
    values = [
        cell
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    ]

    duty_task = next(cell for cell in values if cell.value == "T1.1盤點門市耗材")
    assert duty_task.column == 2
    unassigned_task = next(cell for cell in values if cell.value == "彙整營運週報")
    assert unassigned_task.column == 2
    output_cell = next(
        cell
        for cell in values
        if cell.value == "O1.1.1每週彙整\nO1.1.2每日盤點"
    )
    assert output_cell.data_type == "s"
    knowledge_cell = next(cell for cell in values if cell.value == "K01=1+1")
    assert knowledge_cell.data_type == "s"
    assert "S01試算表操作" in {cell.value for cell in values}
    assert "A01細心" in {cell.value for cell in values}
    workbook.close()
