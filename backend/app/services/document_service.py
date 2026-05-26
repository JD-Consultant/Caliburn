"""
Document export service.

Supports four formats:
  - JSON  : raw OCS document (no file generated, returned as dict)
  - DOCX  : python-docx
  - PDF   : reportlab
  - XLSX  : openpyxl (iCAP-style tabular layout)

All formats derive from profile_data["graph_state"]["ocs_document"].
"""
import json
import logging
import os
import platform
from pathlib import Path
from uuid import UUID

from app.config import settings

logger = logging.getLogger("jobintel")

OUTPUT_DIR = Path(settings.document_output_dir)
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


def _find_cjk_font() -> str | None:
    """Resolve a usable CJK font path: FONT_PATH setting → platform defaults."""
    if settings.font_path:
        if os.path.exists(settings.font_path):
            return settings.font_path
        logger.warning("FONT_PATH '%s' not found, falling back to auto-detect", settings.font_path)

    system = platform.system()
    candidates: list[str] = {
        "Windows": [
            "C:/Windows/Fonts/NotoSansTC-VF.ttf",
            "C:/Windows/Fonts/NotoSansCJKtc-Regular.otf",
            "C:/Windows/Fonts/msjh.ttc",
            "C:/Windows/Fonts/msgothic.ttc",
        ],
        "Darwin": [
            "/Library/Fonts/NotoSansTC-Regular.ttf",
            "/System/Library/Fonts/PingFang.ttc",
            "/Library/Fonts/Arial Unicode MS.ttf",
        ],
    }.get(system, [
        # Linux / Docker (fonts-noto-cjk)
        "/usr/share/fonts/opentype/noto/NotoSansCJKtc-Regular.otf",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/noto-cjk/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/truetype/arphic/uming.ttc",
    ])

    for path in candidates:
        if os.path.exists(path):
            return path

    logger.warning(
        "No CJK font found — Chinese text in PDF will render as boxes. "
        "Set FONT_PATH env var, or on Debian/Ubuntu run: "
        "apt-get install fonts-noto-cjk"
    )
    return None


def _resolve_xlsx_font() -> str:
    """Return a CJK-capable font name for XLSX based on the current platform."""
    return {
        "Windows": "微軟正黑體",
        "Darwin":  "PingFang TC",
    }.get(platform.system(), "Noto Sans CJK TC")


_XLSX_FONT = _resolve_xlsx_font()

SOURCE_LABEL = {"icap_official": "iCAP 標準", "company_defined": "企業自訂"}

ICAP_CONFIDENCE_LABEL = {
    "high": "高信心",
    "medium": "中信心",
    "low": "低信心",
}


def _icap_confidence_label(candidate: dict) -> str:
    return candidate.get("confidence_label") or ICAP_CONFIDENCE_LABEL.get(
        candidate.get("confidence", ""),
        "未評估",
    )


def _icap_recommendation_label(candidate: dict) -> str:
    return {
        "建議參考": "主要參考",
        "部分參考": "可參考",
        "主要參考": "主要參考",
        "可參考": "可參考",
        "低信心": "低信心",
    }.get(candidate.get("recommendation", ""), candidate.get("recommendation", ""))


# display_label → (hex R, hex G, hex B) for DOCX/PDF coloring
_LABEL_RGB = {
    "[訪談確認]": (0x05, 0x96, 0x6B),  # green
    "[iCAP參考]": (0x1D, 0x4E, 0xD8),  # blue
    "[AI整理]":   (0x6B, 0x72, 0x80),  # gray
    "[待確認]":   (0xD9, 0x77, 0x06),  # amber
}
_LABEL_HEX = {
    "[訪談確認]": "#05966B",
    "[iCAP參考]": "#1D4ED8",
    "[AI整理]":   "#6B7280",
    "[待確認]":   "#D97706",
}


# ── Helpers ────────────────────────────────────────────────────────────────────

def _get_ocs(profile_data: dict) -> dict:
    state = profile_data.get("graph_state") or {}
    return state.get("ocs_document") or {}


def _iter_tasks(ocs: dict):
    """Yield (unit, task, task_code, task_name, block) tuples."""
    for unit in ocs.get("ocs_content", {}).get("ocu_units", []):
        for task in unit.get("tasks", []):
            tc = task.get("task_codes", [{}])[0]
            for block in task.get("competency_blocks", []):
                yield unit, task, tc.get("code", ""), tc.get("name", ""), block


# ══════════════════════════════════════════════════════════════════════════════
# JSON
# ══════════════════════════════════════════════════════════════════════════════

def get_ocs_json(profile_data: dict) -> dict:
    """Return the OCS document as a plain dict (no file written)."""
    return _get_ocs(profile_data)


def get_enriched_export_json(profile_data: dict) -> dict:
    """Return enriched JSON export: ocs_document + evidence_refs + icap_reference_pack + quality_scores."""
    ocs   = _get_ocs(profile_data)
    state = profile_data.get("graph_state") or {}

    # ── Flatten all evidence_refs from ocs_document for top-level audit trail ──
    evidence_refs: list[dict] = []
    for unit in ocs.get("ocs_content", {}).get("ocu_units", []):
        for task in unit.get("tasks", []):
            tc         = task.get("task_codes", [{}])[0]
            task_code  = tc.get("code", "")
            task_name  = tc.get("name", "")
            for ref in task.get("evidence_refs", []):
                evidence_refs.append({"task_code": task_code, "task_name": task_name, **ref})
            for block in task.get("competency_blocks", []):
                for ind in block.get("indicators", []):
                    for ref in ind.get("evidence_refs", []):
                        evidence_refs.append({
                            "indicator_code": ind["code"],
                            "task_code": task_code,
                            **ref,
                        })

    # ── iCAP reference pack ────────────────────────────────────────────────────
    icap_reference_pack = {
        "icap_mode": state.get("icap_mode", "company_defined"),
        "icap_hit":  state.get("icap_hit", False),
        "candidates": state.get("icap_candidates", []),
    }

    # ── Per-task quality scores from behavior_indicators ───────────────────────
    quality_scores = {
        bi["task_name"]: {
            "quality_score":  bi.get("quality_score"),
            "quality_status": bi.get("quality_status"),
        }
        for bi in state.get("behavior_indicators", [])
        if bi.get("task_name")
    }

    return {
        "ocs_document":       ocs,
        "evidence_refs":      evidence_refs,
        "icap_reference_pack": icap_reference_pack,
        "quality_scores":     quality_scores,
    }


# ══════════════════════════════════════════════════════════════════════════════
# DOCX
# ══════════════════════════════════════════════════════════════════════════════

def generate_docx(profile_id: UUID, profile_data: dict) -> Path:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    ocs = _get_ocs(profile_data)
    profile = ocs.get("ocs_profile", {})
    doc = Document()

    # ── Title ──
    title = doc.add_heading(profile_data.get("job_title", "職務說明書"), 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = []
    if profile_data.get("department"):
        meta.append(profile_data["department"])
    if profile.get("ocs_code"):
        meta.append(f"代碼：{profile['ocs_code']}")
    if meta:
        p = doc.add_paragraph("　".join(meta))
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    doc.add_paragraph()

    # ── Job description ──
    if profile_data.get("job_summary"):
        doc.add_heading("職務摘要", 1)
        doc.add_paragraph(profile_data["job_summary"])

    # ── iCAP candidates ──
    state = profile_data.get("graph_state") or {}
    candidates = state.get("icap_candidates", [])
    if candidates:
        doc.add_heading("iCAP 職能基準對應", 1)
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text = "職業名稱", "匹配信心", "參考方式"
        for c in candidates:
            row = table.add_row().cells
            row[0].text = c.get("icap_title", "")
            row[1].text = _icap_confidence_label(c)
            row[2].text = _icap_recommendation_label(c)

    # ── OCU units ──
    for unit in ocs.get("ocs_content", {}).get("ocu_units", []):
        doc.add_heading(f"{unit['ocu_code']} {unit['ocu_name']}", 1)

        for task in unit.get("tasks", []):
            tc = task.get("task_codes", [{}])[0]
            doc.add_heading(f"{tc.get('code', '')} {tc.get('name', '')}", 2)

            for block in task.get("competency_blocks", []):
                level = block.get("competency_level", 3)
                p = doc.add_paragraph(f"職能等級：{level}")
                p.runs[0].font.size = Pt(9)
                p.runs[0].font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

                if block.get("indicators"):
                    doc.add_paragraph("行為指標", style="Intense Quote")
                    for ind in block["indicators"]:
                        p = doc.add_paragraph(style="List Bullet")
                        p.add_run(f"[{ind['code']}] ").bold = True
                        lbl = ind.get("display_label", "")
                        if lbl:
                            r, g, b = _LABEL_RGB.get(lbl, (0x6B, 0x72, 0x80))
                            tag = p.add_run(f"{lbl} ")
                            tag.font.size = Pt(8)
                            tag.font.color.rgb = RGBColor(r, g, b)
                        p.add_run(ind["text"])

                if block.get("outputs"):
                    doc.add_paragraph("工作產出", style="Intense Quote")
                    for out in block["outputs"]:
                        doc.add_paragraph(f"[{out['code']}] {out['name']}", style="List Bullet")

                if block.get("knowledge"):
                    doc.add_paragraph("知識 Knowledge", style="Intense Quote")
                    for k in block["knowledge"]:
                        lbl = k.get("display_label") or SOURCE_LABEL.get(k.get("source_type", ""), "")
                        p = doc.add_paragraph(style="List Bullet")
                        p.add_run(f"[{k['code']}] {k['name']}")
                        tag = p.add_run(f"  [{lbl}]")
                        tag.font.size = Pt(8)
                        r, g, b = _LABEL_RGB.get(lbl, (0x6B, 0x72, 0x80))
                        tag.font.color.rgb = RGBColor(r, g, b)

                if block.get("skills"):
                    doc.add_paragraph("技能 Skill", style="Intense Quote")
                    for s in block["skills"]:
                        lbl = s.get("display_label") or SOURCE_LABEL.get(s.get("source_type", ""), "")
                        p = doc.add_paragraph(style="List Bullet")
                        p.add_run(f"[{s['code']}] {s['name']}")
                        tag = p.add_run(f"  [{lbl}]")
                        tag.font.size = Pt(8)
                        r, g, b = _LABEL_RGB.get(lbl, (0x6B, 0x72, 0x80))
                        tag.font.color.rgb = RGBColor(r, g, b)

    # ── Attitudes ──
    attitudes = ocs.get("ocs_attitude", {}).get("attitudes", [])
    if attitudes:
        doc.add_heading("工作態度 Attitude", 1)
        for a in attitudes:
            lbl = a.get("display_label") or SOURCE_LABEL.get(a.get("source_type", ""), "")
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(f"[{a['code']}] {a['name']}")
            tag = p.add_run(f"  [{lbl}]")
            tag.font.size = Pt(8)
            r, g, b = _LABEL_RGB.get(lbl, (0x6B, 0x72, 0x80))
            tag.font.color.rgb = RGBColor(r, g, b)

    # ── 說明與補充事項 ──
    notes = (ocs.get("ocs_profile") or {}).get("notes", "")
    if notes:
        doc.add_heading("說明與補充事項", 1)
        for line in notes.splitlines():
            line = line.strip()
            if line:
                doc.add_paragraph(line)

    out_path = OUTPUT_DIR / f"{profile_id}.docx"
    doc.save(str(out_path))
    logger.info("docx generated: %s", out_path)
    return out_path


# ══════════════════════════════════════════════════════════════════════════════
# PDF (reportlab)
# ══════════════════════════════════════════════════════════════════════════════

def generate_pdf(profile_id: UUID, profile_data: dict) -> Path:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    font_path = _find_cjk_font()
    if font_path:
        pdfmetrics.registerFont(TTFont("CJKFont", font_path))
        fn = "CJKFont"
    else:
        fn = "Helvetica"

    def S(name, **kw):
        kw.setdefault("fontName", fn)
        return ParagraphStyle(name, **kw)

    s_title  = S("T",  fontSize=22, leading=28, spaceAfter=4,  alignment=1,  textColor=colors.HexColor("#111827"))
    s_sub    = S("Su", fontSize=11, leading=16, spaceAfter=12, alignment=1,  textColor=colors.HexColor("#6B7280"))
    s_h1     = S("H1", fontSize=14, leading=20, spaceBefore=12, spaceAfter=5, textColor=colors.HexColor("#1D4ED8"))
    s_h2     = S("H2", fontSize=12, leading=17, spaceBefore=8,  spaceAfter=4, textColor=colors.HexColor("#374151"))
    s_h3     = S("H3", fontSize=10, leading=15, spaceBefore=5,  spaceAfter=3, textColor=colors.HexColor("#1D4ED8"))
    s_body   = S("B",  fontSize=10, leading=15, spaceAfter=3,  textColor=colors.HexColor("#374151"))
    s_meta   = S("M",  fontSize=9,  leading=13, spaceAfter=2,  textColor=colors.HexColor("#9CA3AF"))
    s_bullet = S("Bu", fontSize=10, leading=15, spaceAfter=3,  leftIndent=12,
                 firstLineIndent=-8, textColor=colors.HexColor("#374151"))

    ocs = _get_ocs(profile_data)
    profile = ocs.get("ocs_profile", {})
    story = []

    # Title
    story.append(Paragraph(profile_data.get("job_title", "職務說明書"), s_title))
    dept = profile_data.get("department", "")
    ocs_code = profile.get("ocs_code", "")
    story.append(Paragraph(f"{dept}　{ocs_code}" if dept else ocs_code, s_sub))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#E5E7EB")))
    story.append(Spacer(1, 5 * mm))

    if profile_data.get("job_summary"):
        story.append(Paragraph("職務摘要", s_h1))
        story.append(Paragraph(profile_data["job_summary"], s_body))
        story.append(Spacer(1, 3 * mm))

    # iCAP candidates
    state = profile_data.get("graph_state") or {}
    candidates = state.get("icap_candidates", [])
    if candidates:
        story.append(Paragraph("iCAP 職能基準對應", s_h1))
        tdata = [["職業名稱", "匹配信心", "參考方式"]]
        for c in candidates:
            tdata.append([
                Paragraph(c.get("icap_title", ""), s_body),
                _icap_confidence_label(c),
                _icap_recommendation_label(c),
            ])
        t = Table(tdata, colWidths=[90 * mm, 25 * mm, 45 * mm])
        t.setStyle(TableStyle([
            ("BACKGROUND",  (0, 0), (-1, 0), colors.HexColor("#EFF6FF")),
            ("TEXTCOLOR",   (0, 0), (-1, 0), colors.HexColor("#1D4ED8")),
            ("FONTNAME",    (0, 0), (-1, -1), fn),
            ("FONTSIZE",    (0, 0), (-1, -1), 9),
            ("GRID",        (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E7EB")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F9FAFB")]),
            ("VALIGN",      (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING",  (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(t)
        story.append(Spacer(1, 4 * mm))

    # OCU units
    for unit in ocs.get("ocs_content", {}).get("ocu_units", []):
        story.append(Paragraph(f"{unit['ocu_code']} {unit['ocu_name']}", s_h1))

        for task in unit.get("tasks", []):
            tc = task.get("task_codes", [{}])[0]
            story.append(Paragraph(f"{tc.get('code', '')} {tc.get('name', '')}", s_h2))

            for block in task.get("competency_blocks", []):
                story.append(Paragraph(f"職能等級 {block.get('competency_level', 3)}", s_meta))

                for ind in block.get("indicators", []):
                    lbl = ind.get("display_label", "")
                    lbl_color = _LABEL_HEX.get(lbl, "#6B7280")
                    lbl_html = f' <font color="{lbl_color}" size="8">{lbl}</font>' if lbl else ""
                    story.append(Paragraph(
                        f"<b>[{ind['code']}]</b> 行為指標{lbl_html}", s_h3))
                    story.append(Paragraph(ind["text"], s_body))

                if block.get("outputs"):
                    story.append(Paragraph("工作產出", s_h3))
                    for out in block["outputs"]:
                        story.append(Paragraph(
                            f"• <b>[{out['code']}]</b> {out['name']}", s_bullet))

                if block.get("knowledge"):
                    story.append(Paragraph("知識", s_h3))
                    for k in block["knowledge"]:
                        lbl = k.get("display_label") or SOURCE_LABEL.get(k.get("source_type", ""), "")
                        lbl_color = _LABEL_HEX.get(lbl, "#9CA3AF")
                        story.append(Paragraph(
                            f"• <b>[{k['code']}]</b> {k['name']}  "
                            f"<font color='{lbl_color}' size='8'>[{lbl}]</font>",
                            s_bullet))

                if block.get("skills"):
                    story.append(Paragraph("技能", s_h3))
                    for s in block["skills"]:
                        lbl = s.get("display_label") or SOURCE_LABEL.get(s.get("source_type", ""), "")
                        lbl_color = _LABEL_HEX.get(lbl, "#9CA3AF")
                        story.append(Paragraph(
                            f"• <b>[{s['code']}]</b> {s['name']}  "
                            f"<font color='{lbl_color}' size='8'>[{lbl}]</font>",
                            s_bullet))

            story.append(Spacer(1, 3 * mm))

    # Attitudes
    attitudes = ocs.get("ocs_attitude", {}).get("attitudes", [])
    if attitudes:
        story.append(Paragraph("工作態度", s_h1))
        for a in attitudes:
            lbl = a.get("display_label") or SOURCE_LABEL.get(a.get("source_type", ""), "")
            lbl_color = _LABEL_HEX.get(lbl, "#9CA3AF")
            story.append(Paragraph(
                f"• <b>[{a['code']}]</b> {a['name']}  "
                f"<font color='{lbl_color}' size='8'>[{lbl}]</font>",
                s_bullet))

    # 說明與補充事項
    notes = (ocs.get("ocs_profile") or {}).get("notes", "")
    if notes:
        story.append(Spacer(1, 4 * mm))
        story.append(Paragraph("說明與補充事項", s_h1))
        for line in notes.splitlines():
            line = line.strip()
            if line:
                story.append(Paragraph(line, s_body))

    out_path = OUTPUT_DIR / f"{profile_id}.pdf"
    doc = SimpleDocTemplate(
        str(out_path), pagesize=A4,
        leftMargin=20 * mm, rightMargin=20 * mm,
        topMargin=20 * mm, bottomMargin=20 * mm,
    )
    doc.build(story)
    logger.info("pdf generated: %s", out_path)
    return out_path


# ══════════════════════════════════════════════════════════════════════════════
# XLSX (openpyxl — single-sheet iCAP official layout)
#
# Layout (one sheet "職能基準"):
#   Rows 1-3 : version/profile header block
#   Row  4   : 職業代碼 / 職類別 / 產業別
#   Row  5   : blank
#   Row  6   : main table column headers
#   Row  7+  : one data row per task (P/O/K/S stacked within cells)
#              主要職責 column is vertically merged across tasks of same unit
#   After tasks: Attitude section header + rows
#   After attitudes: 說明與補充事項 placeholder
# ══════════════════════════════════════════════════════════════════════════════

def generate_xlsx(profile_id: UUID, profile_data: dict) -> Path:
    from datetime import date as _date
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "職能基準"

    # ── Styles ──────────────────────────────────────────────────────────────
    _BLUE       = PatternFill("solid", fgColor="2563EB")
    _NAVY       = PatternFill("solid", fgColor="1E3A5F")
    _LIGHT      = PatternFill("solid", fgColor="EFF6FF")
    _GRAY       = PatternFill("solid", fgColor="F3F4F6")
    _AMBER      = PatternFill("solid", fgColor="92400E")
    _NO_FILL    = PatternFill()
    _thin       = Side(style="thin",   color="D1D5DB")
    _med        = Side(style="medium", color="9CA3AF")
    BORDER      = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
    BORDER_MED  = Border(left=_med,  right=_med,  top=_med,  bottom=_med)
    WRAP_TOP    = Alignment(wrap_text=True, vertical="top")
    CENTER      = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def F(bold=False, size=9, color="000000"):
        return Font(name=_XLSX_FONT, bold=bold, size=size, color=color)

    def _cell(row, col, value="", fill=None, font=None, align=None, border=BORDER):
        c = ws.cell(row, col, value)
        if fill:   c.fill   = fill
        if font:   c.font   = font
        if align:  c.alignment = align
        if border: c.border = border
        return c

    def _merge(r1, c1, r2, c2):
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)

    # ── Column widths (7 columns matching official layout) ───────────────────
    # 主要職責 | 工作任務 | 工作產出 | 行為指標 | 職能等級 | K知識 | S技能
    COL_W = [20, 24, 30, 52, 8, 36, 36]
    for i, w in enumerate(COL_W, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ocs         = _get_ocs(profile_data)
    profile_meta = ocs.get("ocs_profile", {})

    # ── Rows 1-3: header block ────────────────────────────────────────────────
    today = _date.today().strftime("%Y/%m/%d")
    _HDR_PAIRS = [
        # (label_col, label_text, value_col, value, value_end_col)
        (1, "版號",       2, "V1",                               2),
        (3, "職能基準代碼", 4, profile_meta.get("ocs_code", ""),  5),
        (6, "職能基準名稱", 7, profile_data.get("job_title", ""), 7),
    ]
    for label_col, label, val_col, val, val_end in _HDR_PAIRS:
        _cell(1, label_col, label, fill=_LIGHT, font=F(bold=True, size=9), align=CENTER)
        _cell(1, val_col,   val,   font=F(size=9), align=CENTER)
        if val_end > val_col:
            _merge(1, val_col, 1, val_end)

    _HDR_PAIRS2 = [
        (1, "部門",   2, profile_data.get("department", ""),          3),
        (4, "職能等級", 5, str(profile_meta.get("ocs_level", 3)),     5),
        (6, "更新版本", 7, today,                                      7),
    ]
    for label_col, label, val_col, val, val_end in _HDR_PAIRS2:
        _cell(2, label_col, label, fill=_LIGHT, font=F(bold=True, size=9), align=CENTER)
        _cell(2, val_col,   val,   font=F(size=9), align=CENTER)
        if val_end > val_col:
            _merge(2, val_col, 2, val_end)

    ws.row_dimensions[3].height = 40
    _cell(3, 1, "職務描述", fill=_LIGHT, font=F(bold=True, size=9), align=CENTER)
    _cell(3, 2, profile_data.get("job_summary", ""), font=F(size=9),
          align=Alignment(wrap_text=True, vertical="center"))
    _merge(3, 2, 3, 7)

    for r in (1, 2, 3):
        ws.row_dimensions[r].height = 18

    # ── Row 4: 職業代碼 / 職類別 / 產業別 ──────────────────────────────────────
    ws.row_dimensions[4].height = 18
    category    = profile_meta.get("category", {})
    occ_code    = category.get("occ_code", "")
    job_cats    = "、".join(c["name"] for c in category.get("job_categories", []))
    industries  = "、".join(i["name"] for i in category.get("industries", []))
    _HDR_PAIRS4 = [
        (1, "職業代碼", 2, occ_code,   3),
        (4, "職類別",   5, job_cats,   5),
        (6, "產業別",   7, industries, 7),
    ]
    for label_col, label, val_col, val, val_end in _HDR_PAIRS4:
        _cell(4, label_col, label, fill=_LIGHT, font=F(bold=True, size=9), align=CENTER)
        _cell(4, val_col,   val,   font=F(size=9), align=CENTER)
        if val_end > val_col:
            _merge(4, val_col, 4, val_end)

    # ── Row 5: blank separator ────────────────────────────────────────────────
    ws.row_dimensions[5].height = 6

    # ── Row 6: main table column headers ─────────────────────────────────────
    MAIN_HDRS = [
        "主要職責",
        "工作任務",
        "工作產出",
        "行為指標",
        "職能\n等級",
        "職能內容\n（K=knowledge 知識）",
        "職能內容\n（S=skills 技能）",
    ]
    ws.row_dimensions[6].height = 36
    for col, hdr in enumerate(MAIN_HDRS, 1):
        _cell(6, col, hdr, fill=_BLUE, font=F(bold=True, size=10, color="FFFFFF"), align=CENTER)

    # ── Rows 7+: one row per task ─────────────────────────────────────────────
    units   = ocs.get("ocs_content", {}).get("ocu_units", [])
    row     = 7

    for unit in units:
        tasks          = unit.get("tasks", [])
        unit_start_row = row

        for task in tasks:
            tc    = task.get("task_codes", [{}])[0]
            block = task.get("competency_blocks", [{}])[0]

            outputs_txt = "\n".join(
                f"{o['code']} {o['name']}" for o in block.get("outputs", [])
            )
            inds_txt = "\n".join(
                f"{i['code']} {i['text']}" for i in block.get("indicators", [])
            )
            k_txt = "\n".join(
                f"{k['code']} {k['name']}  {k.get('display_label', '')}"
                for k in block.get("knowledge", [])
            )
            s_txt = "\n".join(
                f"{s['code']} {s['name']}  {s.get('display_label', '')}"
                for s in block.get("skills", [])
            )

            n_lines = max(
                len(block.get("outputs", [])),
                len(block.get("indicators", [])),
                len(block.get("knowledge", [])),
                len(block.get("skills", [])),
                1,
            )
            ws.row_dimensions[row].height = max(20, n_lines * 16)

            row_fill = _GRAY if (row - 6) % 2 == 1 else _NO_FILL

            # Col 1: 主要職責 — placeholder; merged + filled after all tasks
            _cell(row, 1, f"{unit['ocu_code']}\n{unit['ocu_name']}",
                  font=F(bold=True, size=9), align=CENTER)

            for col, (val, al) in enumerate([
                (f"{tc.get('code','')}\n{tc.get('name','')}", WRAP_TOP),
                (outputs_txt,  WRAP_TOP),
                (inds_txt,     WRAP_TOP),
                (block.get("competency_level", 3), CENTER),
                (k_txt,        WRAP_TOP),
                (s_txt,        WRAP_TOP),
            ], 2):
                c = _cell(row, col, val, font=F(size=9), align=al)
                if row_fill.fill_type:
                    c.fill = row_fill

            row += 1

        # Merge 主要職責 cell across all task rows of this unit
        if len(tasks) > 1:
            _merge(unit_start_row, 1, row - 1, 1)
        ws.cell(unit_start_row, 1).fill = _LIGHT

    # ── Attitude section ──────────────────────────────────────────────────────
    row += 1  # blank gap
    ws.row_dimensions[row].height = 22
    _cell(row, 1, "職能內容（A=attitude 態度）",
          fill=_AMBER, font=F(bold=True, size=10, color="FFFFFF"), align=CENTER)
    _merge(row, 1, row, 7)
    row += 1

    # Attitude sub-headers
    ws.row_dimensions[row].height = 18
    for col, hdr in enumerate(["代碼", "態度名稱", "", "", "iCAP 參照", "來源", ""], 1):
        _cell(row, col, hdr, fill=_BLUE, font=F(bold=True, size=9, color="FFFFFF"), align=CENTER)
    _merge(row, 2, row, 4)
    _merge(row, 6, row, 7)
    row += 1

    for a in ocs.get("ocs_attitude", {}).get("attitudes", []):
        ws.row_dimensions[row].height = 18
        afill = _LIGHT if a.get("icap_ref") else _NO_FILL
        _cell(row, 1, a["code"],   font=F(size=9), align=CENTER, fill=afill)
        c = _cell(row, 2, a["name"], font=F(size=9), align=WRAP_TOP, fill=afill)
        _merge(row, 2, row, 4)
        _cell(row, 5, a.get("icap_ref", ""), font=F(size=9), align=CENTER, fill=afill)
        c = _cell(row, 6,
                  a.get("display_label") or SOURCE_LABEL.get(a.get("source_type", ""), ""),
                  font=F(size=9), align=CENTER, fill=afill)
        _merge(row, 6, row, 7)
        row += 1

    # ── 說明與補充事項 ─────────────────────────────────────────────────────────
    row += 1
    ws.row_dimensions[row].height = 22
    _cell(row, 1, "說明與補充事項",
          fill=_GRAY, font=F(bold=True, size=10), align=CENTER)
    _merge(row, 1, row, 7)
    row += 1

    notes = ocs.get("ocs_profile", {}).get("notes", "")
    ws.row_dimensions[row].height = 72
    _cell(row, 1, notes or "（如有建議擔任此職業之學歷、經歷或其他補充說明，請填寫於此。）",
          font=F(size=9),
          align=Alignment(wrap_text=True, vertical="top"),
          border=Border(left=_thin, right=_thin, top=_thin, bottom=_thin))
    _merge(row, 1, row, 7)

    # ── Page setup (landscape A3 for wide table) ──────────────────────────────
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize   = 8   # A3
    ws.page_setup.fitToPage   = True
    ws.page_setup.fitToWidth  = 1
    ws.page_setup.fitToHeight = 0

    out_path = OUTPUT_DIR / f"{profile_id}.xlsx"
    wb.save(str(out_path))
    logger.info("xlsx generated: %s", out_path)
    return out_path
