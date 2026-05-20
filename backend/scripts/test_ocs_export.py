"""
Standalone test: generate all 4 export formats from a mock OCS document.
Run from backend root: python scripts/test_ocs_export.py
"""
import json
import sys
import uuid
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

MOCK_PROFILE_DATA = {
    "job_title": "行銷企劃專員",
    "department": "行銷部",
    "job_summary": "負責品牌推廣活動策劃、數位內容製作及媒體投放管理，確保品牌曝光與活動效果達標。",
    "graph_state": {
        "icap_candidates": [
            {
                "icap_id": "abc123",
                "ocs_code": "MKT2501-001",
                "icap_title": "行銷企劃人員",
                "similarity": 0.72,
                "recommendation": "建議參考",
                "industries": ["零售業", "服務業"],
                "job_categories": ["行銷企劃"],
            }
        ],
        "ocs_document": {
            "version_info": {
                "versions": [{"status": "最新版本", "ocs_code": "ENT-MKT-001"}]
            },
            "ocs_profile": {
                "ocs_code": "ENT-MKT-001",
                "ocs_name": {"occupation_name": "行銷企劃專員"},
                "job_description": "負責品牌推廣活動策劃、數位內容製作及媒體投放管理。",
                "ocs_level": 3,
                "category": {
                    "job_categories": [{"name": "行銷企劃"}],
                    "industries": [{"name": "零售業"}],
                },
            },
            "ocs_content": {
                "ocu_units": [
                    {
                        "ocu_code": "ENT-U001",
                        "ocu_name": "品牌活動規劃與執行",
                        "tasks": [
                            {
                                "task_codes": [{"code": "ENT-T001", "name": "策劃品牌推廣活動"}],
                                "competency_blocks": [
                                    {
                                        "competency_level": 3,
                                        "indicators": [
                                            {
                                                "code": "ENT-P001",
                                                "text": "在接收行銷指令後，為了提升品牌曝光，與設計部門及外部廠商協作，使用 Notion 及 Meta Ads Manager，完成活動企劃書並執行媒體投放，達成曝光人次 10 萬以上的目標。",
                                            }
                                        ],
                                        "outputs": [
                                            {"code": "ENT-O001", "name": "活動企劃書"},
                                            {"code": "ENT-O002", "name": "媒體投放報告"},
                                        ],
                                        "knowledge": [
                                            {
                                                "code": "ENT-K001",
                                                "name": "品牌行銷基本知識",
                                                "source_type": "icap_official",
                                                "icap_ref": "MKT-K003",
                                            },
                                            {
                                                "code": "ENT-K002",
                                                "name": "公司內部審核與簽核流程",
                                                "source_type": "company_defined",
                                            },
                                        ],
                                        "skills": [
                                            {
                                                "code": "ENT-S001",
                                                "name": "活動企劃與專案管理",
                                                "source_type": "icap_official",
                                                "icap_ref": "MKT-S002",
                                            },
                                            {
                                                "code": "ENT-S002",
                                                "name": "Meta Ads Manager 操作",
                                                "source_type": "company_defined",
                                            },
                                        ],
                                    }
                                ],
                            },
                            {
                                "task_codes": [{"code": "ENT-T002", "name": "數位內容製作與排程"}],
                                "competency_blocks": [
                                    {
                                        "competency_level": 3,
                                        "indicators": [
                                            {
                                                "code": "ENT-P002",
                                                "text": "依據月度行銷日曆，使用 Canva 及 Buffer 製作社群貼文並排程發布，確保每週至少 5 篇貼文按時上線。",
                                            }
                                        ],
                                        "outputs": [
                                            {"code": "ENT-O003", "name": "月度社群內容日曆"},
                                        ],
                                        "knowledge": [
                                            {
                                                "code": "ENT-K003",
                                                "name": "社群媒體演算法知識",
                                                "source_type": "icap_official",
                                                "icap_ref": "MKT-K007",
                                            },
                                        ],
                                        "skills": [
                                            {
                                                "code": "ENT-S003",
                                                "name": "Canva 視覺設計",
                                                "source_type": "company_defined",
                                            },
                                        ],
                                    }
                                ],
                            },
                        ],
                    },
                    {
                        "ocu_code": "ENT-U002",
                        "ocu_name": "行銷數據分析與報告",
                        "tasks": [
                            {
                                "task_codes": [{"code": "ENT-T003", "name": "廣告成效分析與報告"}],
                                "competency_blocks": [
                                    {
                                        "competency_level": 3,
                                        "indicators": [
                                            {
                                                "code": "ENT-P003",
                                                "text": "每月結算後，使用 Google Analytics 及 Meta Insights 彙整廣告成效數據，製作月報並提出次月優化建議，確保 ROAS 達 2.5 以上。",
                                            }
                                        ],
                                        "outputs": [
                                            {"code": "ENT-O004", "name": "月度廣告成效報告"},
                                            {"code": "ENT-O005", "name": "次月優化建議書"},
                                        ],
                                        "knowledge": [
                                            {
                                                "code": "ENT-K004",
                                                "name": "數位廣告指標解讀（CTR/ROAS/CPA）",
                                                "source_type": "icap_official",
                                                "icap_ref": "MKT-K012",
                                            },
                                        ],
                                        "skills": [
                                            {
                                                "code": "ENT-S004",
                                                "name": "Google Analytics 4 操作",
                                                "source_type": "company_defined",
                                            },
                                            {
                                                "code": "ENT-S005",
                                                "name": "數據視覺化報告製作",
                                                "source_type": "icap_official",
                                                "icap_ref": "MKT-S009",
                                            },
                                        ],
                                    }
                                ],
                            }
                        ],
                    },
                ],
            },
            "ocs_attitude": {
                "attitudes": [
                    {
                        "code": "ENT-A001",
                        "name": "積極主動尋求改善行銷效益",
                        "source_type": "icap_official",
                        "icap_ref": "MKT-A002",
                    },
                    {
                        "code": "ENT-A002",
                        "name": "對品牌視覺細節具備高度敏感度",
                        "source_type": "company_defined",
                    },
                    {
                        "code": "ENT-A003",
                        "name": "跨部門溝通時保持開放與彈性",
                        "source_type": "icap_official",
                        "icap_ref": "MKT-A005",
                    },
                ]
            },
        },
    },
}


def test_json():
    from app.services.document_service import get_ocs_json
    ocs = get_ocs_json(MOCK_PROFILE_DATA)
    assert ocs.get("ocs_profile", {}).get("ocs_code") == "ENT-MKT-001"
    units = ocs.get("ocs_content", {}).get("ocu_units", [])
    assert len(units) == 2
    total_tasks = sum(len(u["tasks"]) for u in units)
    assert total_tasks == 3
    print(f"  JSON OK  — {len(units)} OCU units, {total_tasks} tasks")
    return ocs


def test_docx(profile_id):
    from app.services.document_service import generate_docx
    path = generate_docx(profile_id, MOCK_PROFILE_DATA)
    size = path.stat().st_size
    assert size > 5000, f"DOCX too small: {size} bytes"
    print(f"  DOCX OK  — {path.name}  ({size:,} bytes)")


def test_pdf(profile_id):
    from app.services.document_service import generate_pdf
    path = generate_pdf(profile_id, MOCK_PROFILE_DATA)
    size = path.stat().st_size
    assert size > 3000, f"PDF too small: {size} bytes"
    print(f"  PDF  OK  — {path.name}  ({size:,} bytes)")


def test_xlsx(profile_id):
    from app.services.document_service import generate_xlsx
    path = generate_xlsx(profile_id, MOCK_PROFILE_DATA)
    size = path.stat().st_size
    assert size > 3000, f"XLSX too small: {size} bytes"

    # Verify sheet structure
    from openpyxl import load_workbook
    wb = load_workbook(path)
    assert "職能基準" in wb.sheetnames
    assert "知識技能" in wb.sheetnames
    assert "工作態度" in wb.sheetnames
    ws1 = wb["職能基準"]
    ws2 = wb["知識技能"]
    ws3 = wb["工作態度"]
    data_rows_main = ws1.max_row - 8   # 8 header rows
    data_rows_ks   = ws2.max_row - 1
    data_rows_att  = ws3.max_row - 1
    print(f"  XLSX OK  — {path.name}  ({size:,} bytes)")
    print(f"           職能基準: {data_rows_main} data rows | 知識技能: {data_rows_ks} rows | 工作態度: {data_rows_att} rows")


def main():
    profile_id = uuid.uuid4()
    print(f"\nTest profile_id: {profile_id}\n")

    errors = []
    for name, fn, args in [
        ("JSON",  test_json,  []),
        ("DOCX",  test_docx, [profile_id]),
        ("PDF",   test_pdf,  [profile_id]),
        ("XLSX",  test_xlsx, [profile_id]),
    ]:
        try:
            fn(*args)
        except Exception as e:
            print(f"  {name} FAIL — {e}")
            errors.append(name)

    print()
    if errors:
        print(f"FAILED: {errors}")
        sys.exit(1)
    else:
        print("All 4 export formats passed.")


if __name__ == "__main__":
    main()
